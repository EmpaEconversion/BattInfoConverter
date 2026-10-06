"""Unit test auxiliary functions."""

from datetime import datetime

import pytest
from conftest import EXCEL_PATHS, normalize_jsonld
from openpyxl import Workbook

from battinfoconverter_backend.auxiliary import add_to_structure, coerce_date_to_iso
from battinfoconverter_backend.excel_tools import ExcelContainer
from battinfoconverter_backend.json_convert import convert_excel_to_jsonld
from battinfoconverter_backend.registry import Registry


def test_date_coercion() -> None:
    """Check that date coercion works as expected."""
    assert coerce_date_to_iso("2026/01/31") == "2026-01-31"
    assert coerce_date_to_iso("2026-01-31") == "2026-01-31"
    assert coerce_date_to_iso("2026.01.31") == "2026-01-31"
    assert coerce_date_to_iso("31/01/2026") == "2026-01-31"
    assert coerce_date_to_iso("31-01-2026") == "2026-01-31"
    assert coerce_date_to_iso("31.01.2026") == "2026-01-31"

    assert coerce_date_to_iso("   2026/01/31    ") == "2026-01-31"

    assert coerce_date_to_iso(datetime.fromisoformat("2026-01-31")) == "2026-01-31"


def _build(rows: list[tuple[str, float | str, str, str]]) -> dict:
    """Convert (link, value, unit, metadata) rows into a JSON-LD fragment."""
    container = Registry(ExcelContainer(EXCEL_PATHS["flowcell"]))
    jsonld: dict = {}
    for link, value, unit, metadata in rows:
        add_to_structure(jsonld, link.split("-"), value, unit, container, metadata=metadata)
    return jsonld


def test_comment_targets_measured_property_by_type() -> None:
    """A 'type|<Class>' comment attaches to the quantity of that type, not a nested key."""
    jsonld = _build(
        [
            ("hasCatholyte-hasMeasuredProperty-Volume", 30, "mL", "Volume"),
            ("hasCatholyte-hasMeasuredProperty-VolumeFlowRate", 60, "mL/min", "Flow rate"),
            (
                "hasCatholyte-hasMeasuredProperty-type|VolumeFlowRate-rdfs:comment",
                "calibrated",
                "No Unit",
                "Flow rate cal",
            ),
        ]
    )
    volume, flow_rate = jsonld["hasCatholyte"]["hasMeasuredProperty"]
    assert "rdfs:comment" not in volume
    assert "VolumeFlowRate" not in flow_rate
    assert flow_rate["rdfs:comment"] == "Flow rate cal: calibrated"


def test_comment_distinguishes_duplicate_measured_properties() -> None:
    """With two quantities of the same type, the metadata label picks the target."""
    rows = [
        ("hasPositiveElectrode-hasSubstrate-hasMeasuredProperty-Thickness", 560, "um", "Substrate thickness"),
        (
            "hasPositiveElectrode-hasSubstrate-hasMeasuredProperty-Thickness",
            450,
            "um",
            "Thickness in cell under compression",
        ),
    ]
    link = "hasPositiveElectrode-hasSubstrate-hasMeasuredProperty-type|Thickness-rdfs:comment"

    jsonld = _build([*rows, (link, "measured under load", "No Unit", "Thickness in cell comment")])
    uncompressed, compressed = jsonld["hasPositiveElectrode"]["hasSubstrate"]["hasMeasuredProperty"]
    assert "rdfs:comment" not in uncompressed
    assert compressed["rdfs:comment"] == "Thickness in cell comment: measured under load"

    jsonld = _build([*rows, (link, "as delivered", "No Unit", "Substrate thickness comment")])
    uncompressed, compressed = jsonld["hasPositiveElectrode"]["hasSubstrate"]["hasMeasuredProperty"]
    assert uncompressed["rdfs:comment"] == "Substrate thickness comment: as delivered"
    assert "rdfs:comment" not in compressed


# (unit cell, value, ontology link) and whether a missing-unit warning is wanted
UNIT_CASES = [
    ("mm", 12, "hasNegativeElectrode-hasMeasuredProperty-Thickness", False),
    ("No Unit", 12, "hasNegativeElectrode-hasMeasuredProperty-Thickness", True),
    (None, 12, "hasNegativeElectrode-hasMeasuredProperty-Thickness", True),
    ("", 12, "hasNegativeElectrode-hasMeasuredProperty-Thickness", True),
    (None, 7.4, "hasElectrolyte-hasMeasuredProperty-pH", True),
    ("No Unit", "Graphite", "hasNegativeElectrode-hasActiveMaterial", False),
    (None, "Graphite", "hasNegativeElectrode-hasActiveMaterial", False),
    (None, 12345, "hasNegativeElectrode-schema:productID", False),
    (None, "comment|12345", "hasNegativeElectrode-hasActiveMaterial", False),
]


def _unit_workbook(unit: object, value: object, link: str) -> Workbook:
    """A minimal workbook with one row, to exercise the Unit column."""
    wb = Workbook()
    ws = wb.active
    ws.title = "@Schema"
    ws.append(["Metadata", "Value", "Unit", "Priority", "Note", "Ontology link"])
    ws.append(["A row", value, unit, "optional", "", link])
    ws = wb.create_sheet("@Context")
    ws.append(["Term", "IRI"])
    ws.append(["unit", "https://qudt.org/vocab/unit/"])
    ws = wb.create_sheet("@Predicates")
    ws.append(["Predicate", "Default class"])
    for predicate, default in (
        ("hasNegativeElectrode", "NegativeElectrode"),
        ("hasElectrolyte", "Electrolyte"),
        ("hasActiveMaterial", "ActiveMaterial"),
        ("hasMeasuredProperty", None),
    ):
        ws.append([predicate, default])
    ws = wb.create_sheet("@Classes")
    ws.append(["Class", "Note"])
    ws.append(["Graphite", None])
    ws = wb.create_sheet("@Units")
    ws.append(["Unit", "Unit class"])
    ws.append(["mm", "MilliMetre"])
    ws.append(["dimensionless", "unit:UNITLESS"])
    return wb


@pytest.mark.parametrize(("unit", "value", "link", "wants_warning"), UNIT_CASES)
def test_missing_unit_warns_only_for_quantities(
    unit: object, value: object, link: str, wants_warning: bool, caplog: pytest.LogCaptureFixture
) -> None:
    """A number with no unit should be flagged, a name or a literal should not."""
    convert_excel_to_jsonld(_unit_workbook(unit, value, link), validate=False)
    warned = "with no unit, so it is recorded as a comment" in caplog.text
    assert warned == wants_warning


def test_blank_unit_matches_no_unit() -> None:
    """A blank Unit cell and the legacy "No Unit" label must convert the same way."""
    link = "hasNegativeElectrode-hasActiveMaterial"
    blank = convert_excel_to_jsonld(_unit_workbook(None, "Graphite", link), validate=False)
    legacy = convert_excel_to_jsonld(_unit_workbook("No Unit", "Graphite", link), validate=False)
    assert normalize_jsonld(blank) == normalize_jsonld(legacy)


def test_dimensionless_unit_is_a_measurement() -> None:
    """A quantity declared dimensionless is still a measured property, not a comment."""
    out = convert_excel_to_jsonld(
        _unit_workbook("dimensionless", 7.4, "hasElectrolyte-hasMeasuredProperty-pH"), validate=False
    )
    measured = out["hasElectrolyte"]["hasMeasuredProperty"]
    assert measured["@type"] == "pH"
    assert measured["hasNumericalPart"]["hasNumberValue"] == 7.4
    assert measured["hasMeasurementUnit"] == "unit:UNITLESS"
