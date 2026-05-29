"""Test module for backend behaviours."""

import io
from pathlib import Path

import pytest
from conftest import CellFixtures, normalize_jsonld
from openpyxl import load_workbook

from battinfoconverter_backend.json_convert import convert_excel_to_jsonld
from battinfoconverter_backend.templates.template_conversion import (
    dict_to_workbook,
)
from battinfoconverter_backend.validate import validate_jsonld


def test_standard_battinfo_hardcoded_header(tmpdir: Path, coincell: CellFixtures) -> None:
    """Check that coin cell Excel conversion matches expected JSON-LD output when using hardcoded header.

    Required for backwards compatibility.
    """
    new_excel = tmpdir / "NotOntologizeTest.xlsx"
    values_to_update = {
        "Cell type": "NotOntologize",
        "Cell ID": "NotOntologize",
        "Date of cell assembly": "NotOntologize",
        "Institution/company": "NotOntologize",
        "Scientist/technician/operator": "NotOntologize",
        "Project": "Comment",
        "Assembled manually or by robot": "Comment",
        "Schema name": "Comment",
        "Schema version": "Comment",
    }
    wb = load_workbook(coincell.excel)
    sheet = wb["@Schema"]
    for row in sheet.iter_rows():
        col_a = row[0].value
        if col_a in values_to_update:
            row[4].value = values_to_update[col_a]
    wb.save(new_excel)  # Overwrites in place, or use a new name
    converted = convert_excel_to_jsonld(new_excel, debug_mode=False, validate=False)
    assert normalize_jsonld(converted) == normalize_jsonld(coincell.jsonld)


def test_conversion_different_inputs(schema: CellFixtures) -> None:
    """Users should be able to read files in different ways."""
    # pathlib.Path object
    res1 = convert_excel_to_jsonld(schema.excel, validate=False)

    # String object
    res2 = convert_excel_to_jsonld(str(schema.excel), validate=False)

    # Buffered reader object
    with schema.excel.open("rb") as f:
        excel_bytesio = io.BytesIO(f.read())
        res3 = convert_excel_to_jsonld(f, validate=False)

    # Bytes IO object
    res4 = convert_excel_to_jsonld(excel_bytesio, validate=False)

    # Already loaded workbook
    res5 = convert_excel_to_jsonld(load_workbook(schema.excel), validate=False)

    # Should not affect the results
    assert res1 == res2 == res3 == res4 == res5


def test_against_cached_context(coincell: CellFixtures) -> None:
    """Make sure all terms are mapped in the cached context."""
    # The standard filled excel template must pass
    converted = convert_excel_to_jsonld(coincell.excel, validate=True)
    validate_jsonld(converted, errors="raise")

    # Validation should not modify original dict
    assert converted == convert_excel_to_jsonld(coincell.excel, validate=True)

    # Sanity check - these should all fail
    bad_jsonld = converted.copy()
    bad_jsonld["hasSomethingNotAllowed"] = {"@id": "CoinCell"}
    with pytest.raises(ValueError, match="'hasSomethingNotAllowed' was not found"):
        validate_jsonld(bad_jsonld, errors="raise")

    bad_jsonld = converted.copy()
    bad_jsonld["hasComponent"] = {"@id": "ThisDoesNotExist"}
    with pytest.raises(ValueError, match="'ThisDoesNotExist' was not found"):
        validate_jsonld(bad_jsonld, errors="raise")

    bad_jsonld = converted.copy()
    bad_jsonld["hasComponent"] = {"@type": ["CoinCell", "ThisDoesNotExist"]}
    with pytest.raises(ValueError, match="'ThisDoesNotExist' was not found"):
        validate_jsonld(bad_jsonld, errors="raise")

    bad_jsonld = converted.copy()
    bad_jsonld["hasComponent"] = "ThisDoesNotExist"
    with pytest.raises(ValueError, match="'ThisDoesNotExist' was not found"):
        validate_jsonld(bad_jsonld, errors="raise")


def test_bad_jsonld_context(caplog: pytest.LogCaptureFixture) -> None:
    """Check if expected validation warnings/errors trigger."""
    doc = {
        "@context": [
            "https://w3id.org/emmo/domain/battery/context",
            "https://w3id.org/emmo/domain/electrochemistry/context",
        ],
        "@type": "CoinCell",
    }
    with pytest.raises(ValueError, match="There are multiple 'default' vocabularies, you are only allowed one"):
        validate_jsonld(doc, errors="raise")

    doc = {
        "@context": {
            "battery": "https://w3id.org/emmo/domain/batterie",
        },
        "@type": "CoinCell",
    }
    with pytest.raises(ValueError, match=r"Maybe you meant (https://w3id.org/emmo/domain/battery#)?"):
        validate_jsonld(doc, errors="raise")

    caplog.clear()
    doc = {
        "@context": {
            "missing": "https://w3id.org/emmo/domain/somethingwrong",
        },
        "@type": "CoinCell",
        "hasComponent": {
            "@id": "missing:StuffThatCannotBeFound",
        },
    }
    validate_jsonld(doc, errors="warn")
    assert "The URL for 'missing' (https://w3id.org/emmo/domain/somethingwrong) is not a known namespace" in caplog.text
    assert "'CoinCell' has no prefix, but there is no default namespace" in caplog.text
    assert "Term 'missing:StuffThatCannotBeFound' was not found because 'missing' is empty" in caplog.text


def test_bad_prefixed_unit(coincell: CellFixtures, caplog: pytest.LogCaptureFixture) -> None:
    """Check that unit missing from prefixed namespace warns."""
    template = coincell.template.copy()
    template["@Units"]["data"].append({"Item": "foo", "Key": "unit:thisDoesNotExist"})
    template["@Schema"]["data"]["Positive electrode (cathode when battery is discharged)"]["rows"].append(
        {
            "Metadata": "Some made up quantity with a unit missing an IRI",
            "Value": 1.2345,
            "Unit": "foo",
            "Priority": "recommended",
            "Ontology link": "hasPositiveElectrode-hasCurrentCollector-hasMeasuredProperty-Density",
            "Comment": None,
        },
    )
    wb = dict_to_workbook(template)
    convert_excel_to_jsonld(wb, validate=True)
    assert "Term 'unit:thisDoesNotExist' was not found in 'unit'" in caplog.text


def test_bad_default_unit(coincell: CellFixtures, caplog: pytest.LogCaptureFixture) -> None:
    """Check that unit missing from default namespace warns."""
    template = coincell.template.copy()
    template["@Units"]["data"].append({"Item": "foo", "Key": "thisDoesNotExist"})
    template["@Schema"]["data"]["Positive electrode (cathode when battery is discharged)"]["rows"].append(
        {
            "Metadata": "Some made up quantity with a unit missing an IRI",
            "Value": 1.2345,
            "Unit": "foo",
            "Priority": "recommended",
            "Ontology link": "hasPositiveElectrode-hasCurrentCollector-hasMeasuredProperty-Density",
            "Comment": None,
        },
    )
    wb = dict_to_workbook(template)
    convert_excel_to_jsonld(wb, validate=True)
    assert "Term 'thisDoesNotExist' was not found in the default namespace" in caplog.text


def test_missing_unit(coincell: CellFixtures) -> None:
    """Check that missing unit in @Units tab errors."""
    template = coincell.template.copy()
    template["@Schema"]["data"]["Positive electrode (cathode when battery is discharged)"]["rows"].append(
        {
            "Metadata": "Some made up quantity with a unit missing an IRI",
            "Value": 1.2345,
            "Unit": "foo",
            "Priority": "recommended",
            "Ontology link": "hasPositiveElectrode-hasCurrentCollector-hasMeasuredProperty-Density",
            "Comment": None,
        },
    )
    wb = dict_to_workbook(template)
    with pytest.raises(ValueError, match=r"The unit 'foo' was not found in the @Units tab."):
        convert_excel_to_jsonld(wb, validate=True)
