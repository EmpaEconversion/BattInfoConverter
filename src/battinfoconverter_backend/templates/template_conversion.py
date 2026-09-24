"""Functions for converting between the XLSX and JSON template formats.

The JSON template of the xlsx is for internal use only.

Storing as JSON allows us to easily track changes, diff, and version the schema.

It also ensures the Excel files don't become messy over time, and that the
filled, unfilled, and pytest versions of the xlsx files are consistent.
"""

import argparse
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

TEMPLATES_DIR = Path(__file__).parent

COINCELL_TEMPLATE_PATH = TEMPLATES_DIR / "coincell.json"
FLOWCELL_TEMPLATE_PATH = TEMPLATES_DIR / "flowcell.json"
ELECTROLYSIS_TEMPLATE_PATH = TEMPLATES_DIR / "electrolysis.json"

# Excel theme colors to color keys
THEME_INDEX_TO_NAME = {
    1: "grey",
    2: "grey",
    3: "green",
    4: "blue",
    5: "orange",
    6: "grey",
}

# Header and row colors for sections
COLORS = {
    "blue": {
        "header": "FF8EA9DB",
        "rows": [
            None,
            "FFD9E1F2",
        ],
    },
    "grey": {
        "header": "FFA5A5A5",
        "rows": [
            None,
            "FFD0CECE",
        ],
    },
    "green": {
        "header": "FF92D050",
        "rows": [
            None,
            "FFE2EFDA",
        ],
    },
    "orange": {
        "header": "FFFF7815",
        "rows": [
            None,
            "FFFCE4D6",
        ],
    },
    "red": {
        "header": "FFDA9694",
        "rows": [
            None,
            "FFF2DCDB",
        ],
    },
    "purple": {
        "header": "FFB1A0C7",
        "rows": [
            None,
            "FFE4DFEC",
        ],
    },
    "cyan": {
        "header": "FF92CDDC",
        "rows": [
            None,
            "FFDAEEF3",
        ],
    },
}
HEX_TO_NAME = {v["header"]: k for k, v in COLORS.items()}

# Default column widths in 'characters'
COLUMN_WIDTHS = {
    "Metadata": 74,
    "Value": 42,
    "Unit": 9,
    "Priority": 14,
    "Ontology link": 160,
    "Comment": 74,
    "Item": 35,
    "Key": 50,
    "ID": 50,
    "Note": 74,
}

# These sheets are treated as sectioned, others are simple tables
SECTIONED_SHEETS = {"@Schema"}

# These sheets are treated as keyed rows (label + variable-length list of values per row)
KEYED_ROWS_SHEETS = {"@References"}

# The values in these rows are kept even when --empty or empty=True is used
ROWS_TO_KEEP = {
    "Cell type",
    "Schema name",
    "Schema version",
}

# Rows starting with this label switch the references on, and are off in empty templates
INCLUDE_ROW_PREFIX = "Include references"
INCLUDE_ROW_DEFAULT = "no"


def _serialize(value: datetime | date | str | float | None) -> str | float | None:
    """Convert dates to JSON serializable strings."""
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime):
        return value.isoformat()
    return value


def style_header(cell: Cell) -> None:
    """Apply header-style to cell."""
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(vertical="center")


def _is_section_header(row_cells: tuple[Cell]) -> str | None:
    """Check whether a row is a section header."""
    has_values = [bool(c.value) for c in row_cells]
    if (
        has_values[:4] == [True, False, False, False]
        and row_cells[0].font
        and row_cells[0].font.b  # headings are bold
        and not row_cells[0].font.i  # but NOT italic
    ):
        return str(row_cells[0].value)
    return None


def _is_section_subheader(row_cells: tuple[Cell]) -> str | None:
    """Detect a section sub-header."""
    has_values = [bool(c.value) for c in row_cells]
    if (
        has_values[:4] == [True, False, False, False]
        and row_cells[0].font
        and row_cells[0].font.b  # subheadings are bold
        and row_cells[0].font.i  # AND italic
    ):
        return str(row_cells[0].value)
    return None


def _guess_color(row_cells: tuple[Cell]) -> str:
    c = row_cells[0]
    fg = c.fill and c.fill.fgColor
    if fg and fg.type == "rgb":
        hex_val = fg.rgb.upper()
        return HEX_TO_NAME.get(hex_val, "grey")
    if fg and fg.type == "theme":
        return THEME_INDEX_TO_NAME.get(fg.theme, "grey")
    return "grey"


def _get_headers(ws: Worksheet) -> list[str]:
    """Get column headers of sheet as list."""
    return [str(h) for h in next(ws.iter_rows(values_only=True)) if h is not None]


def _read_simple_table(ws: Worksheet) -> list[dict]:
    """Convert a plain sheet into a list of row-dicts.

    Blank rows between entries are kept, as they space out groups of related items,
    but the blank rows trailing the table are dropped - a sheet can claim to be a
    million rows long.
    """
    rows = ws.iter_rows()
    header_row = next(rows, None)
    if header_row is None:
        return []
    headers = [str(c.value) for c in header_row if c.value is not None]
    result = []
    blank_run = 0
    for row in rows:
        values = [_serialize(c.value) for c in row]
        if all(v is None for v in values):
            blank_run += 1
            continue
        result.extend(dict.fromkeys(headers) for _ in range(blank_run))
        blank_run = 0
        result.append(dict(zip(headers, values, strict=False)))
    return result


def _read_sectioned_table(ws: Worksheet) -> dict[str, dict]:
    """Convert a sheet with subsection-header rows into a dict of sections.

    The returned
    """
    sections = {}
    current_section = None
    all_rows = list(ws.iter_rows())

    # First row is always the column-header row
    header_row = all_rows[0]
    headers = [str(c.value) for c in header_row if c.value is not None]

    # Go through each row, if it is a header/subheader, start a new section
    # If it is a normal row, append to the "rows" of that section
    for row in all_rows[1:]:
        if all(c.value is None for c in row):
            continue
        if header := _is_section_header(row):
            current_section = header
            sections[header] = {"type": "section", "color": _guess_color(row), "rows": []}
        elif header := _is_section_subheader(row):
            current_section = header
            sections[header] = {"type": "subsection", "color": _guess_color(row), "rows": []}
        else:
            if current_section is None:
                msg = "All rows must be in a section"
                raise ValueError(msg)
            row_dict = dict(zip(headers, [_serialize(c.value) for c in row], strict=False))
            sections[current_section]["rows"].append(row_dict)
    return sections


def _read_keyed_rows(ws: Worksheet) -> list[dict]:
    """Convert a sheet of label + variable-length value rows into a list of row-dicts.

    Each row is `{"key": <col A>, "values": [remaining cells]}`. A bold row with no
    values is the sheet title (`"type": "header"`) or, when it is filled with a
    colour, a section title (`"type": "section"`) like those of the schema sheet.
    """
    rows = []
    for row in ws.iter_rows():
        if all(c.value is None for c in row):
            continue
        key = _serialize(row[0].value)
        values = [_serialize(c.value) for c in row[1:] if c.value is not None]
        if not values and row[0].font and row[0].font.b:
            if row[0].fill and row[0].fill.patternType:
                rows.append({"key": key, "type": "section", "color": _guess_color(row)})
            else:
                rows.append({"key": key, "type": "header"})
            continue
        rows.append({"key": key, "values": values})
    return rows


def _write_keyed_rows(ws, rows: list[dict], empty: bool = False) -> None:
    """Write a keyed-rows sheet, styled like the sectioned table."""
    n_cols = max((len(entry.get("values", [])) + 1 for entry in rows), default=1)
    ws.column_dimensions["A"].width = COLUMN_WIDTHS["Metadata"]
    for col in range(2, n_cols + 1):
        ws.column_dimensions[get_column_letter(col)].width = COLUMN_WIDTHS["Value"]

    for row_idx, entry in enumerate(rows, 1):
        kind = entry.get("type", "row")
        cell = ws.cell(row=row_idx, column=1, value=entry.get("key"))

        if kind == "header":
            style_header(cell)
            ws.row_dimensions[row_idx].height = 20
            continue

        if kind == "section":
            cell.font = Font(bold=True, size=11)
            cell.alignment = Alignment(vertical="center")
            fill = PatternFill("solid", fgColor=COLORS[entry.get("color", "grey")]["header"])
            for col in range(1, n_cols + 1):
                ws.cell(row=row_idx, column=col).fill = fill
            continue

        values = entry.get("values", [])
        if empty and str(entry.get("key")).startswith(INCLUDE_ROW_PREFIX):
            values = [INCLUDE_ROW_DEFAULT]
        elif empty and entry.get("key") not in ROWS_TO_KEEP:
            values = []
        cell.alignment = Alignment(horizontal="left", vertical="center")
        for col, value in enumerate(values, 2):
            vcell = ws.cell(row=row_idx, column=col, value=value)
            vcell.alignment = Alignment(horizontal="left", vertical="center")


def _write_simple_table(ws, columns: list[str], rows: list[dict]) -> None:
    """Write a simple table."""
    for col_idx, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = col
        style_header(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(col, 42)

    for row_idx, row in enumerate(rows, 2):
        for col, h in enumerate(columns, 1):
            ws.cell(row=row_idx, column=col, value=row.get(h))


def _write_sectioned_table(ws, columns: list[str], sections: dict[str, dict], empty: bool = False) -> None:
    """Write a sectioned table."""
    for col_idx, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = col
        style_header(cell)
        ws.column_dimensions[get_column_letter(col_idx)].width = COLUMN_WIDTHS.get(col, 42)
    ws.row_dimensions[1].height = 20

    n_cols = len(columns)

    current_row = 2

    for title, section in sections.items():
        # Get colors
        color_name = section.get("color", "grey")
        colors = COLORS[color_name]
        header_color = colors["header"]
        row_colors = colors["rows"]

        # Section title
        cell = ws.cell(row=current_row, column=1, value=title)
        cell.font = Font(bold=True, size=11, italic=section.get("type") == "subsection")
        cell.alignment = Alignment(vertical="center")

        # Color title row
        for col in range(1, n_cols + 1):
            ws.cell(row=current_row, column=col).fill = PatternFill("solid", fgColor=header_color)

        current_row += 1

        # Write data rows
        rows = section.get("rows", [])
        for i, row in enumerate(rows):
            color = row_colors[i % len(row_colors)]
            for j, col in enumerate(columns, 1):
                cell = ws.cell(row=current_row, column=j)
                if empty and col == "Value":
                    if row.get("Metadata") in ROWS_TO_KEEP:
                        cell.value = row.get(col)
                else:
                    cell.value = row.get(col)
                if color:
                    cell.fill = PatternFill("solid", fgColor=color)
                cell.alignment = Alignment(horizontal="left", vertical="center")
            current_row += 1


def workbook_to_dict(wb: Workbook) -> dict:
    """Read Excel workbook into dictionary description."""
    output = {}
    for name in wb.sheetnames:
        ws = wb[name]
        if name in SECTIONED_SHEETS:
            output[name] = {"type": "sectioned", "header": _get_headers(ws), "data": _read_sectioned_table(ws)}
        elif name in KEYED_ROWS_SHEETS:
            output[name] = {"type": "keyed_rows", "data": _read_keyed_rows(ws)}
        else:
            output[name] = {"type": "table", "header": _get_headers(ws), "data": _read_simple_table(ws)}
    return output


def xlsx_to_dict(xlsx_path: str | Path) -> dict:
    """Read XLSX into dictionary description."""
    wb = load_workbook(xlsx_path, data_only=True)
    return workbook_to_dict(wb)


def xlsx_to_json(xlsx_path: str | Path, json_path: str | Path) -> None:
    output = xlsx_to_dict(xlsx_path)
    with Path(json_path).open("w", encoding="utf-8") as f:
        f.write(json.dumps(output, indent=2, ensure_ascii=False))


def dict_to_workbook(data: dict, *, empty: bool = False) -> Workbook:
    """Write Excel workbook from dictionary description."""
    wb = Workbook()
    wb.remove(wb.active)  # remove default empty sheet
    for sheet_name, sheet_data in data.items():
        ws = wb.create_sheet(title=sheet_name)
        kind = sheet_data.get("type", "table")
        if kind == "sectioned":
            _write_sectioned_table(ws, sheet_data["header"], sheet_data["data"], empty=empty)
        elif kind == "keyed_rows":
            _write_keyed_rows(ws, sheet_data["data"], empty=empty)
        else:
            _write_simple_table(ws, sheet_data["header"], sheet_data["data"])
    return wb


def dict_to_xlsx(data: dict, xlsx_path: str | Path, *, empty: bool = False) -> None:
    """Write XLSX from dictionary description."""
    wb = dict_to_workbook(data, empty=empty)
    wb.save(xlsx_path)


def json_to_xlsx(json_path: str, xlsx_path: str, *, empty: bool = False) -> None:
    """Read JSON, convert to XLSX."""
    data = json.loads(Path(json_path).read_text(encoding="utf-8"))
    dict_to_xlsx(data, xlsx_path, empty=empty)


def main():
    parser = argparse.ArgumentParser(description="Manage the Excel sheet and its derived formats.")
    sub = parser.add_subparsers(dest="command", required=True)

    p_dump = sub.add_parser("dump", help="Dump XLSX to JSON template.")
    p_dump.add_argument("xlsx", help="Source Excel file")
    p_dump.add_argument("json", help="Output JSON template")

    p_load = sub.add_parser("load", help="Rebuild XLSX from a JSON template.")
    p_load.add_argument("json", help="Source JSON template")
    p_load.add_argument("xlsx", help="Output Excel file")
    p_load.add_argument(
        "--empty",
        action="store_true",
        default=False,
        help="Generate the XLSX with empty value cells (default: filled).",
    )

    args = parser.parse_args()

    if args.command == "dump":
        xlsx_to_json(args.xlsx, args.json)
    elif args.command == "load":
        json_to_xlsx(args.json, args.xlsx, empty=args.empty)


if __name__ == "__main__":
    main()
